"""Dataset export application, capability, and local writer tests."""

import asyncio
import csv
from dataclasses import replace
from datetime import UTC, datetime, timedelta
from pathlib import Path
from typing import Any
from uuid import UUID, uuid4

import pyarrow as pa  # type: ignore[import-untyped]
import pyarrow.parquet as pq  # type: ignore[import-untyped]
import pytest
from openpyxl import load_workbook  # type: ignore[import-untyped]

from datepipette.application.access_control import DatasetAccessAuthorizer
from datepipette.application.exports import (
    DatasetExportService,
    ExportDispatchUnavailableError,
    ExportDownloadService,
    ExportExecutionResult,
    ExportExpiredError,
    ExportFileUnavailableError,
    ExportGenerationTimeoutError,
    ExportLimitExceededError,
    ExportNotFoundError,
    ExportNotReadyError,
    ExportWorkerService,
    InvalidExportDownloadTokenError,
    export_query_document,
    export_query_from_document,
)
from datepipette.application.snapshot_queries import InvalidSnapshotQueryError
from datepipette.domain.access_control import DatasetAccess, DatasetPermission
from datepipette.domain.exports import (
    ExportArtifact,
    ExportClaimOutcome,
    ExportFormat,
    ExportJob,
    ExportRequestOutcome,
    ExportState,
)
from datepipette.domain.identity import AccountStatus, SystemRole, User
from datepipette.domain.snapshot_queries import (
    ActiveSnapshotQueryContext,
    SnapshotFilterOperator,
    SnapshotQuery,
    SnapshotQueryField,
    SnapshotQueryFilter,
    SnapshotQuerySort,
    SnapshotRefreshState,
    SnapshotSortDirection,
)
from datepipette.infrastructure.export_dispatch import CeleryExportDispatcher
from datepipette.infrastructure.export_tokens import EncryptedExportDownloadTokenCodec
from datepipette.infrastructure.local_exports import LocalDatasetExportExecutor, LocalExportStorage
from datepipette.infrastructure.object_storage import (
    LocalObjectStorage,
    ObjectExportDownloadResolver,
    ObjectStorageRegistry,
)

DATASET_ID = uuid4()
SNAPSHOT_ID = uuid4()
ACTOR_ID = uuid4()
ID_FIELD = SnapshotQueryField(
    uuid4(), "record_id", "field_id", "bigint", False, True, "Record ID", None
)
TEXT_FIELD = SnapshotQueryField(uuid4(), "label", "field_text", "text", True, True, "Label", None)
DATE_FIELD = SnapshotQueryField(uuid4(), "day", "field_date", "date", True, True, "Day", None)
FIELDS = (ID_FIELD, TEXT_FIELD, DATE_FIELD)


def _admin(role: SystemRole = SystemRole.ADMIN) -> User:
    now = datetime.now(UTC)
    return User(ACTOR_ID, "admin@example.com", "Admin", role, AccountStatus.ACTIVE, None, now, now)


def _context() -> ActiveSnapshotQueryContext:
    return ActiveSnapshotQueryContext(
        DATASET_ID,
        SNAPSHOT_ID,
        "local",
        "snapshot.parquet",
        datetime.now(UTC),
        3,
        FIELDS,
        SnapshotRefreshState.IDLE,
        None,
        False,
    )


def _query() -> SnapshotQuery:
    return SnapshotQuery(
        (ID_FIELD.field_id, TEXT_FIELD.field_id, DATE_FIELD.field_id),
        (SnapshotQueryFilter(ID_FIELD.field_id, SnapshotFilterOperator.GREATER_THAN, 1),),
        (SnapshotQuerySort(ID_FIELD.field_id, SnapshotSortDirection.DESCENDING),),
        1,
        None,
    )


def _job(**changes: Any) -> ExportJob:
    now = datetime.now(UTC)
    values: dict[str, Any] = {
        "id": uuid4(),
        "dataset_id": DATASET_ID,
        "snapshot_id": SNAPSHOT_ID,
        "requested_by": ACTOR_ID,
        "format": ExportFormat.CSV,
        "state": ExportState.QUEUED,
        "query": _query(),
        "dispatch_attempts": 0,
        "dispatched_at": None,
        "attempt_count": 0,
        "object_key": None,
        "row_count": None,
        "byte_size": None,
        "checksum_sha256": None,
        "content_type": None,
        "error_code": None,
        "requested_at": now,
        "started_at": None,
        "finished_at": None,
        "expires_at": now + timedelta(hours=24),
        "updated_at": now,
    }
    values.update(changes)
    return ExportJob(**values)


class FakeDispatcher:
    def __init__(self, fail: bool = False) -> None:
        self.fail = fail
        self.ids: list[UUID] = []

    def dispatch(self, export_id: UUID) -> None:
        if self.fail:
            raise RuntimeError("broker secret")
        self.ids.append(export_id)


class FakeStore:
    def __init__(self, context: ActiveSnapshotQueryContext | None = None) -> None:
        self.context = context
        self.job: ExportJob | None = None
        self.result: ExportExecutionResult | None = None
        self.claim_outcome = ExportClaimOutcome.CLAIMED

    async def get_active_context(self, dataset_id: UUID) -> ActiveSnapshotQueryContext | None:
        return self.context if dataset_id == DATASET_ID else None

    async def get_dataset_access(self, *, actor_id: UUID, dataset_id: UUID) -> DatasetAccess | None:
        if dataset_id != DATASET_ID:
            return None
        return DatasetAccess(dataset_id, False, frozenset(DatasetPermission))

    async def request(self, **values: Any) -> ExportRequestOutcome | None:
        if self.context is None or values["snapshot_id"] != self.context.snapshot_id:
            return None
        self.job = _job(
            id=values["export_id"],
            format=values["format"],
            query=export_query_from_document(values["query_document"]),
            expires_at=values["expires_at"],
        )
        return ExportRequestOutcome(self.job, False, True)

    async def get(self, *, dataset_id: UUID, export_id: UUID) -> ExportJob | None:
        return (
            self.job if self.job and self.job.id == export_id and dataset_id == DATASET_ID else None
        )

    async def get_by_id(self, export_id: UUID) -> ExportJob | None:
        return self.job if self.job and self.job.id == export_id else None

    async def list(
        self, *, dataset_id: UUID, requested_by: UUID | None, limit: int
    ) -> tuple[ExportJob, ...]:
        return (
            (self.job,)
            if self.job
            and dataset_id == DATASET_ID
            and (requested_by is None or requested_by == self.job.requested_by)
            and limit
            else ()
        )

    async def record_dispatch_attempt(self, export_id: UUID, occurred_at: datetime) -> None:
        assert self.job and self.job.id == export_id
        self.job = replace(
            self.job, dispatch_attempts=self.job.dispatch_attempts + 1, updated_at=occurred_at
        )

    async def mark_dispatched(self, export_id: UUID, occurred_at: datetime) -> None:
        assert self.job and self.job.id == export_id
        self.job = replace(self.job, dispatched_at=occurred_at, updated_at=occurred_at)

    async def claim(
        self, export_id: UUID, occurred_at: datetime, policy_revision: int
    ) -> tuple[ExportClaimOutcome, ExportJob | None]:
        if self.job is None or self.job.id != export_id:
            return ExportClaimOutcome.NOT_FOUND, None
        self.job = replace(self.job, policy_revision=policy_revision)
        return self.claim_outcome, self.job

    async def complete(
        self, export_id: UUID, result: ExportExecutionResult, occurred_at: datetime
    ) -> None:
        assert self.job and self.job.id == export_id
        self.result = result

    async def get_context(self, job: ExportJob) -> ActiveSnapshotQueryContext | None:
        return self.context if job.snapshot_id == SNAPSHOT_ID else None


class FakeExecutor:
    def __init__(self, failure: Exception | None = None) -> None:
        self.failure = failure

    async def execute(self, job: ExportJob, context: ActiveSnapshotQueryContext) -> ExportArtifact:
        if self.failure is not None:
            raise self.failure
        assert job.snapshot_id == context.snapshot_id
        return ExportArtifact("exports/result.csv", 2, 20, "0" * 64, "text/csv")


def test_export_service_validates_dispatches_lists_and_hides_unauthorized() -> None:
    store = FakeStore(_context())
    dispatcher = FakeDispatcher()
    service = DatasetExportService(store, dispatcher, DatasetAccessAuthorizer(store))
    outcome = asyncio.run(
        service.request(
            actor=_admin(),
            dataset_id=DATASET_ID,
            format=ExportFormat.CSV,
            query=_query(),
            idempotency_key="retry-me",
            request_id="request",
        )
    )
    assert dispatcher.ids == [outcome.job.id]
    assert outcome.job.dispatch_attempts == 1
    assert outcome.job.dispatched_at is not None
    assert (
        asyncio.run(service.get(actor=_admin(), dataset_id=DATASET_ID, export_id=outcome.job.id)).id
        == outcome.job.id
    )
    assert asyncio.run(service.list(actor=_admin(), dataset_id=DATASET_ID, limit=10))
    analyst = _admin(SystemRole.ANALYST)
    assert asyncio.run(service.list(actor=analyst, dataset_id=DATASET_ID, limit=10))
    assert (
        asyncio.run(service.get(actor=analyst, dataset_id=DATASET_ID, export_id=outcome.job.id)).id
        == outcome.job.id
    )
    other_analyst = replace(analyst, id=uuid4())
    assert not asyncio.run(service.list(actor=other_analyst, dataset_id=DATASET_ID, limit=10))
    with pytest.raises(ExportNotFoundError):
        asyncio.run(
            service.get(
                actor=other_analyst,
                dataset_id=DATASET_ID,
                export_id=outcome.job.id,
            )
        )

    with pytest.raises(ExportNotFoundError):
        asyncio.run(
            service.get(
                actor=_admin(SystemRole.VIEWER), dataset_id=DATASET_ID, export_id=outcome.job.id
            )
        )
    with pytest.raises(InvalidSnapshotQueryError):
        asyncio.run(
            service.request(
                actor=_admin(),
                dataset_id=DATASET_ID,
                format=ExportFormat.CSV,
                query=replace(_query(), field_ids=(uuid4(),)),
                idempotency_key="bad",
                request_id="request",
            )
        )


def test_export_dispatch_and_worker_failures_are_safe() -> None:
    store = FakeStore(_context())
    service = DatasetExportService(store, FakeDispatcher(fail=True), DatasetAccessAuthorizer(store))
    with pytest.raises(ExportDispatchUnavailableError):
        asyncio.run(
            service.request(
                actor=_admin(),
                dataset_id=DATASET_ID,
                format=ExportFormat.XLSX,
                query=_query(),
                idempotency_key="key",
                request_id="request",
            )
        )
    assert store.job is not None
    worker = ExportWorkerService(store, FakeExecutor(), policy_revision=7)
    assert asyncio.run(worker.execute(store.job.id)) is ExportClaimOutcome.CLAIMED
    assert store.result and store.result.succeeded
    assert store.job.policy_revision == 7
    store.result = None
    assert (
        asyncio.run(
            ExportWorkerService(
                store,
                FakeExecutor(RuntimeError("filesystem secret")),
                policy_revision=8,
            ).execute(store.job.id)
        )
        is ExportClaimOutcome.CLAIMED
    )
    assert store.result and store.result.error_code == "export_generation_failed"
    store.result = None
    asyncio.run(
        ExportWorkerService(
            store,
            FakeExecutor(ExportGenerationTimeoutError()),
            policy_revision=9,
        ).execute(store.job.id)
    )
    assert store.result and store.result.error_code == "export_timeout"
    store.result = None
    asyncio.run(
        ExportWorkerService(
            store,
            FakeExecutor(ExportLimitExceededError()),
            policy_revision=10,
        ).execute(store.job.id)
    )
    assert store.result and store.result.error_code == "export_limit_exceeded"
    store.claim_outcome = ExportClaimOutcome.TERMINAL
    assert asyncio.run(worker.execute(store.job.id)) is ExportClaimOutcome.TERMINAL


def test_query_document_round_trip_and_rejects_corruption() -> None:
    assert export_query_from_document(export_query_document(_query())) == _query()
    with pytest.raises(InvalidSnapshotQueryError):
        export_query_from_document({"fields": ["bad"]})


def _write_snapshot(root: Path) -> None:
    table = pa.table(
        {
            "field_id": pa.array([1, 2, 3], type=pa.int64()),
            "field_text": pa.array(["ordinary", "=SUM(A1:A2)", "+cmd"], type=pa.string()),
            "field_date": pa.array(
                [datetime(2026, 1, 1).date(), datetime(2026, 1, 2).date(), None],
                type=pa.date32(),
            ),
        }
    )
    pq.write_table(table, root / "snapshot.parquet")


@pytest.mark.parametrize("format", [ExportFormat.CSV, ExportFormat.XLSX])
def test_local_export_streams_filtered_sorted_safe_files(
    tmp_path: Path, format: ExportFormat
) -> None:
    _write_snapshot(tmp_path)
    executor = LocalDatasetExportExecutor(
        tmp_path,
        memory_limit_mb=64,
        threads=1,
        temporary_directory_limit_mb=64,
        maximum_rows=10,
        fetch_rows=1,
    )
    job = _job(format=format)
    artifact = asyncio.run(executor.execute(job, _context()))
    assert artifact.row_count == 2
    assert artifact.byte_size > 0
    assert len(artifact.checksum_sha256) == 64
    path = LocalExportStorage(tmp_path).resolve(artifact.object_key)
    if format is ExportFormat.CSV:
        with path.open(encoding="utf-8-sig", newline="") as stream:
            rows = list(csv.reader(stream))
        assert rows[1][0] == "3"
        assert rows[1][1] == "'+cmd"
        assert rows[2][1] == "'=SUM(A1:A2)"
    else:
        workbook = load_workbook(path, read_only=True, data_only=False)
        rows = list(workbook.active.iter_rows(values_only=True))
        workbook.close()
        assert int(rows[1][0]) == 3
        assert rows[1][1] == "'+cmd"
        assert rows[2][1] == "'=SUM(A1:A2)"


def test_local_export_limits_and_storage_paths(tmp_path: Path) -> None:
    _write_snapshot(tmp_path)
    executor = LocalDatasetExportExecutor(
        tmp_path,
        memory_limit_mb=64,
        threads=1,
        temporary_directory_limit_mb=64,
        maximum_rows=1,
        fetch_rows=2,
    )
    job = _job()
    with pytest.raises(ValueError, match="row limit"):
        asyncio.run(executor.execute(job, _context()))
    assert not (tmp_path / f"exports/{DATASET_ID}/{job.id}.csv").exists()
    with pytest.raises(ValueError):
        LocalExportStorage(tmp_path).resolve("../secret")
    with pytest.raises(ValueError):
        LocalExportStorage(tmp_path).resolve("missing.csv")
    with pytest.raises(ValueError):
        asyncio.run(executor.execute(job, replace(_context(), object_key="../snapshot.parquet")))

    byte_limited = LocalDatasetExportExecutor(
        tmp_path,
        memory_limit_mb=64,
        threads=1,
        temporary_directory_limit_mb=64,
        maximum_rows=10,
        maximum_bytes=1,
        fetch_rows=2,
    )
    with pytest.raises(ExportLimitExceededError, match="artifact byte limit"):
        asyncio.run(byte_limited.execute(job, _context()))

    timed_out = LocalDatasetExportExecutor(
        tmp_path,
        memory_limit_mb=64,
        threads=1,
        temporary_directory_limit_mb=64,
        maximum_rows=10,
        timeout_seconds=0,
        fetch_rows=2,
    )
    with pytest.raises(ExportGenerationTimeoutError):
        asyncio.run(timed_out.execute(job, _context()))


def test_download_tokens_and_service_are_reissuable_and_bounded(tmp_path: Path) -> None:
    target = tmp_path / "exports/result.csv"
    target.parent.mkdir(parents=True)
    target.write_text("id\n1\n")
    codec = EncryptedExportDownloadTokenCodec("test-installation-key")
    store = FakeStore(_context())
    store.job = _job(
        state=ExportState.SUCCEEDED,
        object_key="exports/result.csv",
        content_type="text/csv",
        row_count=1,
        byte_size=5,
        checksum_sha256=None,
        storage_backend="local",
    )
    local = LocalObjectStorage(tmp_path)
    service = ExportDownloadService(
        store,
        codec,
        ObjectExportDownloadResolver(ObjectStorageRegistry(local.backend, {local.backend: local})),
        DatasetAccessAuthorizer(store),
    )
    token, expires_at = asyncio.run(
        service.issue(actor=_admin(), dataset_id=DATASET_ID, export_id=store.job.id)
    )
    assert expires_at <= datetime.now(UTC) + timedelta(minutes=5, seconds=1)
    download = asyncio.run(service.resolve(token))
    assert download.path is not None
    assert Path(download.path) == target
    assert download.filename.endswith(".csv")
    middle = len(token) // 2
    tampered = token[:middle] + ("A" if token[middle] != "A" else "B") + token[middle + 1 :]
    with pytest.raises(InvalidExportDownloadTokenError):
        asyncio.run(service.resolve(tampered))

    store.job = replace(store.job, state=ExportState.RUNNING)
    with pytest.raises(ExportNotReadyError):
        asyncio.run(service.issue(actor=_admin(), dataset_id=DATASET_ID, export_id=store.job.id))
    store.job = replace(
        store.job, state=ExportState.SUCCEEDED, expires_at=datetime.now(UTC) - timedelta(seconds=1)
    )
    with pytest.raises(ExportExpiredError):
        asyncio.run(service.issue(actor=_admin(), dataset_id=DATASET_ID, export_id=store.job.id))
    store.job = replace(
        store.job, expires_at=datetime.now(UTC) + timedelta(hours=1), object_key="missing.csv"
    )
    token, _ = asyncio.run(
        service.issue(actor=_admin(), dataset_id=DATASET_ID, export_id=store.job.id)
    )
    with pytest.raises(ExportFileUnavailableError):
        asyncio.run(service.resolve(token))


def test_celery_export_dispatch_contract() -> None:
    class FakeCelery:
        def __init__(self) -> None:
            self.call: dict[str, Any] = {}

        def send_task(self, name: str, **kwargs: Any) -> None:
            self.call = {"name": name, **kwargs}

    fake = FakeCelery()
    export_id = uuid4()
    CeleryExportDispatcher(fake).dispatch(export_id)
    assert fake.call["name"] == "datepipette.export.execute"
    assert fake.call["queue"] == "export"
    assert fake.call["task_id"] == f"export-{export_id}"
